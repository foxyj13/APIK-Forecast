import logging
import os

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import pandas as pd

FT_BOLD = Font(bold=True)
FT_BOLD_BLUE = Font(color="000066CC", bold=True)
COL_WIDHT = 12
COL_WIDHT_2X = 24


class Writer:
    def __init__(self, args: dict):
        self.mode = args["main-args"]["mode"]
        self.start_date_str = args["main-args"]["start-date"]
        self.end_date_str = args["main-args"]["end-date"]
        self.step_date_str = args["main-args"]["step-date"]
        self.output_folder = args["main-args"]["output-folder"]

        self.time_depth = args["forecast-args"]["time-depth"]
        self.time_forecast = args["forecast-args"]["time-forecast"]
        self.forecast_type = args["forecast-args"]["forecast-type"]

        self.exp_name = args["data-args"]["experiment-name"]

        self.metrics_filename_prefix = (
            "apik-forecast_metrics_"
            + self.exp_name
            + "_"
            + self.start_date_str
            + "_"
            + self.end_date_str
            + "_step"
            + self.step_date_str
            + "_"
            + self.forecast_type
            + "_td"
            + self.time_depth
            + "_tf"
            + self.time_forecast
        )

        self.metrics_daily_filename_prefix = (
            "apik-forecast_metrics-daily_"
            + self.exp_name
            + "_"
            + self.start_date_str
            + "_"
            + self.end_date_str
            + "_step"
            + self.step_date_str
            + "_"
            + self.forecast_type
            + "_td"
            + self.time_depth
            + "_tf"
            + self.time_forecast
        )
        self.metrics_daily_sum_filename_prefix = (
            "apik-forecast_metrics-daily-sum_"
            + self.exp_name
            + "_"
            + self.start_date_str
            + "_"
            + self.end_date_str
            + "_step"
            + self.step_date_str
            + "_"
            + self.forecast_type
            + "_td"
            + self.time_depth
            + "_tf"
            + self.time_forecast
        )

        self.data_filename_prefix = (
            "apik-forecast_data_"
            + self.exp_name
            + "_"
            + self.start_date_str
            + "_"
            + self.end_date_str
            + "_step"
            + self.step_date_str
            + "_"
            + self.forecast_type
            + "_td"
            + self.time_depth
            + "_tf"
            + self.time_forecast
        )

        self.self_metrics = ["min", "max", "mean", "med"]
        self.sub_self = [
            "obs_td" + self.time_depth,
            "loc_td" + self.time_depth,
            "glob_td" + self.time_depth,
            "obs_tf" + self.time_forecast,
            "loc_tf" + self.time_forecast,
            "glob_tf" + self.time_forecast,
        ]
        self.mutual_metrics = ["me", "mae", "rmse", "r2", "mre", "nse", "kge"]
        self.sub_mutual = [
            "loc_td" + self.time_depth,
            "glob_td" + self.time_depth,
            "loc_tf" + self.time_forecast,
            "glob_tf" + self.time_forecast,
        ]

        self.name_metrics_from_metrics = ["min", "max", "mean", "std"]

    def _collect_info(self, in_data_0: dict) -> dict:
        logging.info(
            "Сбор информации о расчетах для формирования 1-го листа с аннотацией"
        )

        info = {}

        for idx_station, station in enumerate(in_data_0["stations"]):
            st_code = station["code"]
            info[st_code] = {}
            info[st_code]["station"] = {
                "full_name": station["full_name"],
                "lat": station["lat"],
                "lon": station["lon"],
            }

            info[st_code]["parameters"] = {}
            for idx_par, par_name in enumerate(station["parameters"]):
                info[st_code]["parameters"][par_name] = {
                    "code": station["parameters"][par_name]["code"],
                    "name": station["parameters"][par_name]["full_name"],
                    "min": station["parameters"][par_name]["min_value"],
                    "max": station["parameters"][par_name]["max_value"],
                }

                if "prepast" in station["timeline"] and station["timeline"]["prepast"]:
                    #  DB
                    info[st_code]["parameters"][par_name]["data-source"] = "db"
                    info[st_code]["parameters"][par_name]["predictors"] = station[
                        "parameters"
                    ][par_name]["db_predictors"]
                    info[st_code]["parameters"][par_name]["om_parameter"] = None
                else:
                    # OM
                    info[st_code]["parameters"][par_name]["data-source"] = "om"
                    info[st_code]["parameters"][par_name]["predictors"] = station[
                        "parameters"
                    ][par_name]["om_predictors"]
                    info[st_code]["parameters"][par_name]["om_parameter"] = station[
                        "parameters"
                    ][par_name]["om_parameter"]

        info["config"] = {
            "mode": in_data_0["args"]["mode"],
            "forecast-type": in_data_0["args"]["forecast_type"],
            "time-depth": in_data_0["args"]["time_depth"],
            "time-forecast": in_data_0["args"]["time_forecast"],
        }
        if (
            "prepast" in in_data_0["stations"][0]["timeline"]
            and in_data_0["stations"][0]["timeline"]["prepast"]
        ):
            info["config"]["om_model"] = None
        else:
            info["config"]["om_model"] = in_data_0["config_ini"]["om"]["model"]

        return info

    def _collect_ml_models_info(self, in_data: dict, now_dates: list) -> dict:
        logging.info(
            "Формирование датафреймов с перечнем и характеристиками ML-моделей для записи в файлы"
        )

        dict_ml_models = {}

        for now_date in now_dates:
            dict_ml_models[now_date] = {}
            for station in in_data[now_date]["stations"]:
                station_code = station["code"]
                if station_code not in dict_ml_models[now_date]:
                    dict_ml_models[now_date][station_code] = {}
                for par_name, par_info in station["parameters"].items():
                    # if par_name not in dict_ml_models[now_date][station_code]:
                    #     dict_ml_models[now_date][station_code][par_name] = {}
                    if ("chosen_model" in par_info) and par_info["chosen_model"]:
                        if (
                            par_info["chosen_model"]["model_type"] == "best"
                            or par_info["chosen_model"]["model_type"] == "only"
                        ):
                            model_name = par_info["chosen_model"]["model"]
                            dict_ml_models[now_date][station_code][
                                par_name
                            ] = f"{model_name}({par_info['models'][model_name]['model_cfg']['cfg_pars']})[{par_info['models'][model_name]['model_cfg']['cfg_type']}]"
                        elif par_info["chosen_model"]["model_type"] == "mean":
                            dict_ml_models[now_date][station_code][
                                par_name
                            ] = f"{par_info['chosen_model']['model_type']}"
                    else:
                        dict_ml_models[now_date][station_code][par_name] = None

        return dict_ml_models

    # exp_data =
    # { "stations": {
    #   "code": {"full_name", "lat", "lon",
    #       "parameters": {
    #           <par_name>: {"code", "full_name", "om_parameter",
    #                       "obs": [float, ...],
    #                       "local": {"past": {<now_date>: {"data": [float, ...], "timeline":[datetime, ...]}, ...}
    #                               "future": {<now_date>: {"data": [float, ...], "timeline":[datetime, ...]}, ...}},
    #                       "global": {"past": {<now_date>: {"data": [float, ...], "timeline":[datetime, ...]}, ...}
    #                                "future": {<now_date>: {"data": [float, ...], "timeline":[datetime, ...]}, ...}}

    #           }
    #       }
    #   }, ...
    # },
    #  "obs_timeline": [datetime, ...] # сгенерирован нужный
    # }
    def _gen_dfs_data(self, exp_data: dict, now_dates: list) -> dict:
        def get_df_from_data(
            forec_type, time_type, par_data, now_dates, columns, indeces
        ):
            logging.info(
                "Формирование датафрейма для прогноза %s %s", forec_type, time_type
            )

            df = pd.DataFrame(columns=columns, index=indeces)

            # Заполнение датафрейма
            df.loc[indeces, "obs"] = par_data["obs"]

            for now_date in now_dates:
                idxs = [
                    str(tl_date.replace(tzinfo=None))
                    for tl_date in par_data[forec_type][time_type][now_date]["timeline"]
                ]
                df.loc[idxs, now_date] = par_data[forec_type][time_type][now_date][
                    "data"
                ]

            return df

        logging.info(
            "Формирование датафреймов с данными: 1 датафрейм = 1 станция 1 метео-параметр 1 тип прогноза (локальный/глобальный & вперед/назад) по всей серии прогнозов и наблюдений за период расчета"
        )

        sheet_names = [
            "local_tf" + self.time_forecast,
            "global_tf" + self.time_forecast,
            "local_td" + self.time_depth,
            "global_td" + self.time_depth,
        ]

        columns = ["obs"] + now_dates
        indeces = [
            str(tl_date.replace(tzinfo=None)) for tl_date in exp_data["obs_timeline"]
        ]

        dfs = {}

        for station_code, station_data in exp_data["stations"].items():
            dfs[station_code] = {}
            for par_name, par_data in station_data["parameters"].items():
                logging.info(
                    "Формирование датафреймов с данными прогнозов для станции %s и параметра %s",
                    station_code,
                    par_name,
                )

                dfs[station_code][par_name] = {}
                for sheet_name in sheet_names:
                    if "local" in sheet_name:
                        if "tf" in sheet_name:
                            dfs[station_code][par_name][sheet_name] = get_df_from_data(
                                "local", "future", par_data, now_dates, columns, indeces
                            )
                        elif "td" in sheet_name:
                            dfs[station_code][par_name][sheet_name] = get_df_from_data(
                                "local", "past", par_data, now_dates, columns, indeces
                            )
                    elif "global" in sheet_name:
                        if "tf" in sheet_name:
                            dfs[station_code][par_name][sheet_name] = get_df_from_data(
                                "global",
                                "future",
                                par_data,
                                now_dates,
                                columns,
                                indeces,
                            )
                        elif "td" in sheet_name:
                            dfs[station_code][par_name][sheet_name] = get_df_from_data(
                                "global", "past", par_data, now_dates, columns, indeces
                            )

        return dfs

    def _get_station_codes_par_names(self, metrics) -> dict[str, list[str]]:
        station_pars = {}

        for now_date_str, stations in metrics.items():
            for station_code, parameters in stations.items():
                if station_code not in station_pars:
                    station_pars[station_code] = set()
                station_pars[station_code].update(parameters.keys())

        for station_code, par_names in station_pars.items():
            par_names = sorted(list(par_names))

        return station_pars

    def _gen_dfs_metrics(self, metrics: dict) -> dict:

        # def get_station_codes_par_names() -> dict[str, list[str]]:
        #     station_pars = {}

        #     for now_date_str, stations in metrics.items():
        #         for station_code, parameters in stations.items():
        #             if station_code not in station_pars:
        #                 station_pars[station_code] = set()
        #             station_pars[station_code].update(parameters.keys())

        #     for station_code, par_names in station_pars.items():
        #         par_names = sorted(list(par_names))

        #     return station_pars

        logging.info(
            "Формирование датафреймов со всеми оценками: 1 датафрейм = 1 станция 1 метео-параметр"
        )

        station_parameters = self._get_station_codes_par_names(metrics)

        columns = pd.MultiIndex.from_tuples(
            [(metr, sub) for metr in self.self_metrics for sub in self.sub_self]
            + [(metr, sub) for metr in self.mutual_metrics for sub in self.sub_mutual]
        )

        dfs = {}
        for station_code, par_names in station_parameters.items():
            dfs[station_code] = {}
            for par_name in par_names:

                dfs_st_par = pd.DataFrame(columns=columns)
                # Заполнить датафрейм значениями метрик для каждой даты now_date_str, где индексом будет now_date_str, а значениями будут соответствующие метрики
                for now_date_str in metrics.keys():
                    if (
                        station_code in metrics[now_date_str]
                        and par_name in metrics[now_date_str][station_code]
                    ):
                        row_values = []
                        for metr in self.self_metrics:
                            row_values.append(
                                metrics[now_date_str][station_code][par_name]["obs"][
                                    "past"
                                ][metr]
                            )

                            if metrics[now_date_str][station_code][par_name]["local"][
                                "past"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "local"
                                    ]["past"][metr]
                                )
                            else:
                                row_values.append(None)

                            if (
                                "global"
                                in metrics[now_date_str][station_code][par_name]
                            ) and (
                                metrics[now_date_str][station_code][par_name]["global"][
                                    "past"
                                ]
                            ):
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "global"
                                    ]["past"][metr]
                                )
                            else:
                                row_values.append(None)

                            row_values.append(
                                metrics[now_date_str][station_code][par_name]["obs"][
                                    "future"
                                ][metr]
                            )
                            if metrics[now_date_str][station_code][par_name]["local"][
                                "future"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "local"
                                    ]["future"][metr]
                                )
                            else:
                                row_values.append(None)

                            if (
                                "global"
                                in metrics[now_date_str][station_code][par_name]
                            ) and (
                                metrics[now_date_str][station_code][par_name]["global"][
                                    "future"
                                ]
                            ):
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "global"
                                    ]["future"][metr]
                                )
                            else:
                                row_values.append(None)

                        for metr in self.mutual_metrics:
                            if metrics[now_date_str][station_code][par_name]["local"][
                                "past"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "local"
                                    ]["past"][metr]
                                )
                            else:
                                row_values.append(None)

                            if (
                                "global"
                                in metrics[now_date_str][station_code][par_name]
                            ) and metrics[now_date_str][station_code][par_name][
                                "global"
                            ][
                                "past"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "global"
                                    ]["past"][metr]
                                )
                            else:
                                row_values.append(None)

                            if metrics[now_date_str][station_code][par_name]["local"][
                                "future"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "local"
                                    ]["future"][metr]
                                )
                            else:
                                row_values.append(None)

                            if (
                                "global"
                                in metrics[now_date_str][station_code][par_name]
                            ) and metrics[now_date_str][station_code][par_name][
                                "global"
                            ][
                                "future"
                            ]:
                                row_values.append(
                                    metrics[now_date_str][station_code][par_name][
                                        "global"
                                    ]["future"][metr]
                                )
                            else:
                                row_values.append(None)

                        dfs_st_par.loc[now_date_str] = row_values

                dfs[station_code][par_name] = dfs_st_par

        return dfs

    def _gen_dfs_metrics_daily(self, metrics: dict) -> tuple[dict, dict]:
        logging.info(
            "Формирование датафреймов с посуточными метриками: 1 датафрейм = 1 станция 1 метеопараметр 1 метрика наблюдения vs. прогноз вперед (локальный и глобальный) по всей серии прогнозов за период расчета"
        )

        station_parameters = self._get_station_codes_par_names(metrics)

        sheet_names = [metr for metr in self.mutual_metrics]

        n_days = int(self.time_forecast[:-1]) + 1
        indeces = [f"{idx_d}d" for idx_d in range(n_days)]

        now_dates = list(metrics.keys())
        columns_daily = pd.MultiIndex.from_tuples(
            [
                (forec, now_date)
                for forec in [
                    "local",
                    "global",
                ]
                for now_date in now_dates
            ]
        )

        columns_daily_sum = pd.MultiIndex.from_tuples(
            [
                (forec, metr)
                for forec in [
                    "local",
                    "global",
                ]
                for metr in ["min", "max", "mean", "std"]
            ]
        )

        dfs_daily_dict = {}
        dfs_daily_sum_dict = {}
        for station_code, par_names in station_parameters.items():
            dfs_daily_dict[station_code] = {}
            dfs_daily_sum_dict[station_code] = {}
            for par_name in sorted(par_names):
                dfs_daily_dict[station_code][par_name] = {}
                dfs_daily_sum_dict[station_code][par_name] = {}
                for sheet_name in sheet_names:
                    # dfs_daily_dict[station_code][par_name][sheet_name] = {}

                    dfs_daily = pd.DataFrame(columns=columns_daily, index=indeces)
                    dfs_daily_sum = pd.DataFrame(
                        columns=columns_daily_sum, index=indeces
                    )

                    for forec in ["local", "global"]:
                        for now_date in now_dates:
                            if (
                                sheet_name + "_day"
                                in metrics[now_date][station_code][par_name][forec][
                                    "future"
                                ]
                            ):
                                dfs_daily[(forec, now_date)] = metrics[now_date][
                                    station_code
                                ][par_name][forec]["future"][sheet_name + "_day"]

                    for forec in ["local", "global"]:
                        for idx_day in indeces:
                            for metr in ["min", "max", "mean", "std"]:
                                dfs_daily_sum.loc[idx_day, (forec, metr)] = (
                                    dfs_daily.loc[idx_day, forec].agg([metr]).loc[metr]
                                )

                    dfs_daily_dict[station_code][par_name][sheet_name] = dfs_daily
                    dfs_daily_sum_dict[station_code][par_name][
                        sheet_name
                    ] = dfs_daily_sum

        return dfs_daily_dict, dfs_daily_sum_dict

    def _write_info_sheet_to_station_file(self, wb, info, st_code):
        logging.info(
            "Запись информации о расчетах в файл с метриками для станции %s", st_code
        )

        ws = wb.create_sheet(title="Info", index=0)

        ws.column_dimensions["A"].width = COL_WIDHT
        ws.column_dimensions["B"].width = COL_WIDHT
        ws.column_dimensions["C"].width = COL_WIDHT
        ws.column_dimensions["D"].width = COL_WIDHT
        ws.column_dimensions["E"].width = COL_WIDHT
        ws.column_dimensions["F"].width = COL_WIDHT_2X
        ws.column_dimensions["G"].width = COL_WIDHT_2X

        ws["A1"] = "QC-отчет для периода:"
        ws["A1"].font = FT_BOLD_BLUE

        ws["A2"], ws["B2"], ws["C2"] = "Начало", "Конец", "Шаг"
        ws["A2"].font, ws["B2"].font, ws["C2"].font = FT_BOLD, FT_BOLD, FT_BOLD

        ws["A3"], ws["B3"], ws["C3"] = (
            self.start_date_str,
            self.end_date_str,
            self.step_date_str,
        )

        ws["A5"] = "Станция"
        ws["A5"].font = FT_BOLD_BLUE

        ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Код", "Название", "Широта", "Долгота"
        ws["A6"].font, ws["B6"].font, ws["C6"].font, ws["D6"].font = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        ws["A7"], ws["B7"], ws["C7"], ws["D7"] = (
            st_code,
            info[st_code]["station"]["full_name"],
            info[st_code]["station"]["lat"],
            info[st_code]["station"]["lon"],
        )

        ws["A9"] = "Параметры"
        ws["A9"].font = FT_BOLD_BLUE

        (
            ws["A10"],
            ws["B10"],
            ws["C10"],
            ws["D10"],
            ws["E10"],
            ws["F10"],
            ws["G10"],
        ) = (
            "Код",
            "Суффикс",
            "Название",
            "Минимум",
            "Максимум",
            "Предикторы",
            "ОМ-соответствие",
        )
        (
            ws["A10"].font,
            ws["B10"].font,
            ws["C10"].font,
            ws["D10"].font,
            ws["E10"].font,
            ws["F10"].font,
            ws["G10"].font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        irow = 11
        for par_name, parameter in info[st_code]["parameters"].items():
            ws.cell(column=1, row=irow).value = parameter["code"]  # A
            ws.cell(column=2, row=irow).value = par_name  # B
            ws.cell(column=3, row=irow).value = parameter["name"]  # C
            ws.cell(column=4, row=irow).value = parameter["min"]  # D
            ws.cell(column=5, row=irow).value = parameter["max"]  # E
            ws.cell(column=6, row=irow).value = "; ".join(parameter["predictors"])  # F
            ws.cell(column=7, row=irow).value = parameter["om_parameter"]  # G

            irow = irow + 1

        irow = irow + 1
        ws.cell(column=1, row=irow).value = "Параметры расчета"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE

        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=2, row=irow + 1).value,
            ws.cell(column=3, row=irow + 1).value,
            ws.cell(column=4, row=irow + 1).value,
            ws.cell(column=5, row=irow + 1).value,
        ) = (
            "Режим расчета",
            "Тип прогноза",
            "Источник данных",
            "Период обучения",
            "Период прогноза",
        )
        (
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).font,
            ws.cell(column=3, row=irow + 1).font,
            ws.cell(column=4, row=irow + 1).font,
            ws.cell(column=5, row=irow + 1).font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=2, row=irow + 2).value,
            ws.cell(column=3, row=irow + 2).value,
            ws.cell(column=4, row=irow + 2).value,
            ws.cell(column=5, row=irow + 2).value,
        ) = (
            info["config"]["mode"],
            info["config"]["forecast-type"],
            info["config"]["om_model"],
            info["config"]["time-depth"],
            info["config"]["time-forecast"],
        )

        irow = irow + 4
        ws.cell(column=1, row=irow).value = "Метрики"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE

        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).value,
        ) = ("min", FT_BOLD, "Минимальное значение")
        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=1, row=irow + 2).font,
            ws.cell(column=2, row=irow + 2).value,
        ) = ("max", FT_BOLD, "Максимальное значение")
        (
            ws.cell(column=1, row=irow + 3).value,
            ws.cell(column=1, row=irow + 3).font,
            ws.cell(column=2, row=irow + 3).value,
        ) = ("mean", FT_BOLD, "Среднее значение")
        (
            ws.cell(column=1, row=irow + 4).value,
            ws.cell(column=1, row=irow + 4).font,
            ws.cell(column=2, row=irow + 4).value,
        ) = ("med", FT_BOLD, "Медианное значение")

        (
            ws.cell(column=1, row=irow + 5).value,
            ws.cell(column=1, row=irow + 5).font,
            ws.cell(column=2, row=irow + 5).value,
        ) = ("me", FT_BOLD, "Средняя ошибка")
        (
            ws.cell(column=1, row=irow + 6).value,
            ws.cell(column=1, row=irow + 6).font,
            ws.cell(column=2, row=irow + 6).value,
        ) = ("mae", FT_BOLD, "Средняя абсолютная ошибка")
        (
            ws.cell(column=1, row=irow + 7).value,
            ws.cell(column=1, row=irow + 7).font,
            ws.cell(column=2, row=irow + 7).value,
        ) = ("rmse", FT_BOLD, "Среднеквадратическая ошибка")
        (
            ws.cell(column=1, row=irow + 8).value,
            ws.cell(column=1, row=irow + 8).font,
            ws.cell(column=2, row=irow + 8).value,
        ) = ("r2", FT_BOLD, "Коэффициент детерминации")
        (
            ws.cell(column=1, row=irow + 9).value,
            ws.cell(column=1, row=irow + 9).font,
            ws.cell(column=2, row=irow + 9).value,
        ) = ("mre", FT_BOLD, "Средняя относительная ошибка (она же MAPE)")
        (
            ws.cell(column=1, row=irow + 10).value,
            ws.cell(column=1, row=irow + 10).font,
            ws.cell(column=2, row=irow + 10).value,
        ) = ("nse", FT_BOLD, "Коэффициент эффективности Нэша-Сатклиффа")
        (
            ws.cell(column=1, row=irow + 11).value,
            ws.cell(column=1, row=irow + 11).font,
            ws.cell(column=2, row=irow + 11).value,
        ) = ("kge", FT_BOLD, "Коэффициент эффективности Клинга-Гупты")

        irow = irow + 13
        ws.cell(column=1, row=irow).value = "Данные"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE
        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).value,
        ) = (
            "obs_td<X>d",
            FT_BOLD,
            "Наблюдения из БД ИМКЭС на периоде обучения из <X> дней (time-depth=<X>d)",
        )
        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=1, row=irow + 2).font,
            ws.cell(column=2, row=irow + 2).value,
        ) = (
            "loc_td<X>d",
            FT_BOLD,
            "Локальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
        )
        (
            ws.cell(column=1, row=irow + 3).value,
            ws.cell(column=1, row=irow + 3).font,
            ws.cell(column=2, row=irow + 3).value,
        ) = (
            "glob_td<X>d",
            FT_BOLD,
            "Глобальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
        )

        (
            ws.cell(column=1, row=irow + 4).value,
            ws.cell(column=1, row=irow + 4).font,
            ws.cell(column=2, row=irow + 4).value,
        ) = (
            "obs_tf<Y>d",
            FT_BOLD,
            "Наблюдения из БД ИМКЭС на периоде прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )
        (
            ws.cell(column=1, row=irow + 5).value,
            ws.cell(column=1, row=irow + 5).font,
            ws.cell(column=2, row=irow + 5).value,
        ) = (
            "loc_tf<Y>d",
            FT_BOLD,
            "Локальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )
        (
            ws.cell(column=1, row=irow + 6).value,
            ws.cell(column=1, row=irow + 6).font,
            ws.cell(column=2, row=irow + 6).value,
        ) = (
            "glob_tf<Y>d",
            FT_BOLD,
            "Глобальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )

    def _write_info_sheet_to_parameter_file(self, wb, info, par_name):
        logging.info(
            "Запись информации о расчетах в файл с метриками для параметра %s", par_name
        )

        ws = wb.create_sheet(title="Info", index=0)

        ws.column_dimensions["A"].width = COL_WIDHT
        ws.column_dimensions["B"].width = COL_WIDHT
        ws.column_dimensions["C"].width = COL_WIDHT
        ws.column_dimensions["D"].width = COL_WIDHT
        ws.column_dimensions["E"].width = COL_WIDHT
        ws.column_dimensions["F"].width = COL_WIDHT_2X
        ws.column_dimensions["G"].width = COL_WIDHT_2X

        ws["A1"] = "QC-отчет для периода:"
        ws["A1"].font = FT_BOLD_BLUE

        ws["A2"], ws["B2"], ws["C2"] = "Начало", "Конец", "Шаг"
        ws["A2"].font, ws["B2"].font, ws["C2"].font = FT_BOLD, FT_BOLD, FT_BOLD

        ws["A3"], ws["B3"], ws["C3"] = (
            self.start_date_str,
            self.end_date_str,
            self.step_date_str,
        )

        ws["A5"] = "Параметр"
        ws["A5"].font = FT_BOLD_BLUE

        (
            ws["A6"],
            ws["B6"],
            ws["C6"],
            ws["D6"],
            ws["E6"],
            ws["F6"],
            ws["G6"],
        ) = (
            "Коды",
            "Суффикс",
            "Название",
            "Минимум",
            "Максимум",
            "Предикторы",
            "ОМ-соответствие",
        )
        (
            ws["A6"].font,
            ws["B6"].font,
            ws["C6"].font,
            ws["D6"].font,
            ws["E6"].font,
            ws["F6"].font,
            ws["G6"].font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )
        # Строка 7: заполняется ниже

        ws["A9"] = "Станции"
        ws["A9"].font = FT_BOLD_BLUE

        ws["A10"], ws["B10"], ws["C10"], ws["D10"] = (
            "Код",
            "Название",
            "Широта",
            "Долгота",
        )
        ws["A10"].font, ws["B10"].font, ws["C10"].font, ws["D10"].font = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        irow = 11
        par_codes = set()

        for st_code, station in info.items():
            if st_code != "config":
                if par_name in station["parameters"]:
                    any_st_code = st_code
                    par_codes.add(station["parameters"][par_name]["code"])

                    ws.cell(column=1, row=irow).value = st_code  # A
                    ws.cell(column=2, row=irow).value = station["station"][
                        "full_name"
                    ]  # B
                    ws.cell(column=3, row=irow).value = station["station"]["lat"]  # C
                    ws.cell(column=4, row=irow).value = station["station"]["lon"]  # D

                    irow = irow + 1

        (
            ws["A7"],
            ws["B7"],
            ws["C7"],
            ws["D7"],
            ws["E7"],
            ws["F7"],
            ws["G7"],
        ) = (
            "; ".join(list(par_codes)),
            par_name,
            info[any_st_code]["parameters"][par_name]["name"],
            info[any_st_code]["parameters"][par_name]["min"],
            info[any_st_code]["parameters"][par_name]["max"],
            "; ".join(info[any_st_code]["parameters"][par_name]["predictors"]),
            info[any_st_code]["parameters"][par_name]["om_parameter"],
        )

        irow = irow + 1
        ws.cell(column=1, row=irow).value = "Параметры расчета"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE

        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=2, row=irow + 1).value,
            ws.cell(column=3, row=irow + 1).value,
            ws.cell(column=4, row=irow + 1).value,
            ws.cell(column=5, row=irow + 1).value,
        ) = (
            "Режим расчета",
            "Тип прогноза",
            "Источник данных",
            "Период обучения",
            "Период прогноза",
        )
        (
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).font,
            ws.cell(column=3, row=irow + 1).font,
            ws.cell(column=4, row=irow + 1).font,
            ws.cell(column=5, row=irow + 1).font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=2, row=irow + 2).value,
            ws.cell(column=3, row=irow + 2).value,
            ws.cell(column=4, row=irow + 2).value,
            ws.cell(column=5, row=irow + 2).value,
        ) = (
            info["config"]["mode"],
            info["config"]["forecast-type"],
            info["config"]["om_model"],
            info["config"]["time-depth"],
            info["config"]["time-forecast"],
        )

        irow = irow + 4
        ws.cell(column=1, row=irow).value = "Метрики"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE

        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).value,
        ) = ("min", FT_BOLD, "Минимальное значение")
        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=1, row=irow + 2).font,
            ws.cell(column=2, row=irow + 2).value,
        ) = ("max", FT_BOLD, "Максимальное значение")
        (
            ws.cell(column=1, row=irow + 3).value,
            ws.cell(column=1, row=irow + 3).font,
            ws.cell(column=2, row=irow + 3).value,
        ) = ("mean", FT_BOLD, "Среднее значение")
        (
            ws.cell(column=1, row=irow + 4).value,
            ws.cell(column=1, row=irow + 4).font,
            ws.cell(column=2, row=irow + 4).value,
        ) = ("med", FT_BOLD, "Медианное значение")

        (
            ws.cell(column=1, row=irow + 5).value,
            ws.cell(column=1, row=irow + 5).font,
            ws.cell(column=2, row=irow + 5).value,
        ) = ("me", FT_BOLD, "Средняя ошибка")
        (
            ws.cell(column=1, row=irow + 6).value,
            ws.cell(column=1, row=irow + 6).font,
            ws.cell(column=2, row=irow + 6).value,
        ) = ("mae", FT_BOLD, "Средняя абсолютная ошибка")
        (
            ws.cell(column=1, row=irow + 7).value,
            ws.cell(column=1, row=irow + 7).font,
            ws.cell(column=2, row=irow + 7).value,
        ) = ("rmse", FT_BOLD, "Среднеквадратическая ошибка")
        (
            ws.cell(column=1, row=irow + 8).value,
            ws.cell(column=1, row=irow + 8).font,
            ws.cell(column=2, row=irow + 8).value,
        ) = ("r2", FT_BOLD, "Коэффициент детерминации")
        (
            ws.cell(column=1, row=irow + 9).value,
            ws.cell(column=1, row=irow + 9).font,
            ws.cell(column=2, row=irow + 9).value,
        ) = ("mre", FT_BOLD, "Средняя относительная ошибка (она же MAPE)")
        (
            ws.cell(column=1, row=irow + 10).value,
            ws.cell(column=1, row=irow + 10).font,
            ws.cell(column=2, row=irow + 10).value,
        ) = ("nse", FT_BOLD, "Коэффициент эффективности Нэша-Сатклиффа")
        (
            ws.cell(column=1, row=irow + 11).value,
            ws.cell(column=1, row=irow + 11).font,
            ws.cell(column=2, row=irow + 11).value,
        ) = ("kge", FT_BOLD, "Коэффициент эффективности Клинга-Гупты")

        irow = irow + 13
        ws.cell(column=1, row=irow).value = "Данные"
        ws.cell(column=1, row=irow).font = FT_BOLD_BLUE
        (
            ws.cell(column=1, row=irow + 1).value,
            ws.cell(column=1, row=irow + 1).font,
            ws.cell(column=2, row=irow + 1).value,
        ) = (
            "obs_td<X>d",
            FT_BOLD,
            "Наблюдения из БД ИМКЭС на периоде обучения из <X> дней (time-depth=<X>d)",
        )
        (
            ws.cell(column=1, row=irow + 2).value,
            ws.cell(column=1, row=irow + 2).font,
            ws.cell(column=2, row=irow + 2).value,
        ) = (
            "loc_td<X>d",
            FT_BOLD,
            "Локальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
        )
        (
            ws.cell(column=1, row=irow + 3).value,
            ws.cell(column=1, row=irow + 3).font,
            ws.cell(column=2, row=irow + 3).value,
        ) = (
            "glob_td<X>d",
            FT_BOLD,
            "Глобальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
        )

        (
            ws.cell(column=1, row=irow + 4).value,
            ws.cell(column=1, row=irow + 4).font,
            ws.cell(column=2, row=irow + 4).value,
        ) = (
            "obs_tf<Y>d",
            FT_BOLD,
            "Наблюдения из БД ИМКЭС на периоде прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )
        (
            ws.cell(column=1, row=irow + 5).value,
            ws.cell(column=1, row=irow + 5).font,
            ws.cell(column=2, row=irow + 5).value,
        ) = (
            "loc_tf<Y>d",
            FT_BOLD,
            "Локальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )
        (
            ws.cell(column=1, row=irow + 6).value,
            ws.cell(column=1, row=irow + 6).font,
            ws.cell(column=2, row=irow + 6).value,
        ) = (
            "glob_tf<Y>d",
            FT_BOLD,
            "Глобальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
        )

    def _write_info_sheet_to_expdata_file(
        self, wb, data_type, info, st_code, par_name
    ) -> None:
        logging.info(
            "Запись информации о расчетах в файл с экспортируемыми данными для станции %s и параметра %s",
            st_code,
            par_name,
        )

        ws = wb.create_sheet(title="Info", index=0)

        ws.column_dimensions["A"].width = COL_WIDHT
        ws.column_dimensions["B"].width = COL_WIDHT
        ws.column_dimensions["C"].width = COL_WIDHT
        ws.column_dimensions["D"].width = COL_WIDHT
        ws.column_dimensions["E"].width = COL_WIDHT
        ws.column_dimensions["F"].width = COL_WIDHT_2X
        ws.column_dimensions["G"].width = COL_WIDHT_2X

        if data_type == "data_export":
            ws["A1"] = "Экспорт данных для периода:"
            ws["A1"].font = FT_BOLD_BLUE
        elif data_type == "metrics_daily":
            ws["A1"] = "Посуточные метрики для периода:"
            ws["A1"].font = FT_BOLD_BLUE
        elif data_type == "metrics_daily_sum":
            ws["A1"] = "Обобщенные посуточные метрики для периода:"
            ws["A1"].font = FT_BOLD_BLUE

        ws["A2"], ws["B2"], ws["C2"] = "Начало", "Конец", "Шаг"
        ws["A2"].font, ws["B2"].font, ws["C2"].font = FT_BOLD, FT_BOLD, FT_BOLD

        ws["A3"], ws["B3"], ws["C3"] = (
            self.start_date_str,
            self.end_date_str,
            self.step_date_str,
        )

        ws["A5"] = "Станция"
        ws["A5"].font = FT_BOLD_BLUE

        ws["A6"], ws["B6"], ws["C6"], ws["D6"] = "Код", "Название", "Широта", "Долгота"
        ws["A6"].font, ws["B6"].font, ws["C6"].font, ws["D6"].font = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )

        ws["A7"], ws["B7"], ws["C7"], ws["D7"] = (
            st_code,
            info[st_code]["station"]["full_name"],
            info[st_code]["station"]["lat"],
            info[st_code]["station"]["lon"],
        )

        ws["A9"] = "Параметр"
        ws["A9"].font = FT_BOLD_BLUE

        (
            ws["A10"],
            ws["B10"],
            ws["C10"],
            ws["D10"],
            ws["E10"],
            ws["F10"],
            ws["G10"],
        ) = (
            "Коды",
            "Суффикс",
            "Название",
            "Минимум",
            "Максимум",
            "Предикторы",
            "ОМ-соответствие",
        )
        (
            ws["A10"].font,
            ws["B10"].font,
            ws["C10"].font,
            ws["D10"].font,
            ws["E10"].font,
            ws["F10"].font,
            ws["G10"].font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )
        (
            ws["A11"],
            ws["B11"],
            ws["C11"],
            ws["D11"],
            ws["E11"],
            ws["F11"],
            ws["G11"],
        ) = (
            info[st_code]["parameters"][par_name]["code"],
            par_name,
            info[st_code]["parameters"][par_name]["name"],
            info[st_code]["parameters"][par_name]["min"],
            info[st_code]["parameters"][par_name]["max"],
            "; ".join(info[st_code]["parameters"][par_name]["predictors"]),
            info[st_code]["parameters"][par_name]["om_parameter"],
        )

        ws["A13"] = "Параметры расчета"
        ws["A13"].font = FT_BOLD_BLUE
        (
            ws["A14"],
            ws["B14"],
            ws["C14"],
            ws["D14"],
            ws["E14"],
        ) = (
            "Режим расчета",
            "Тип прогноза",
            "Источник данных",
            "Период обучения",
            "Период прогноза",
        )
        (
            ws["A14"].font,
            ws["B14"].font,
            ws["C14"].font,
            ws["D14"].font,
            ws["E14"].font,
        ) = (
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
            FT_BOLD,
        )
        (
            ws["A15"],
            ws["B15"],
            ws["C15"],
            ws["D15"],
            ws["E15"],
        ) = (
            info["config"]["mode"],
            info["config"]["forecast-type"],
            info["config"]["om_model"],
            info["config"]["time-depth"],
            info["config"]["time-forecast"],
        )

        ws["A17"] = "Данные"
        ws["A17"].font = FT_BOLD_BLUE

        if data_type == "data_export":
            (
                ws["A18"].value,
                ws["A18"].font,
                ws["B18"].value,
            ) = (
                "obs",
                FT_BOLD,
                "Наблюдения из БД ИМКЭС",
            )
            (
                ws["A19"].value,
                ws["A19"].font,
                ws["B19"].value,
            ) = (
                "loc_tf<Y>d",
                FT_BOLD,
                "Локальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
            )
            (
                ws["A20"].value,
                ws["A20"].font,
                ws["B20"].value,
            ) = (
                "glob_tf<Y>d",
                FT_BOLD,
                "Глобальный прогноз вперед от текущей даты из <Y> дней (time-forecast=<Y>d)",
            )
            (
                ws["A21"].value,
                ws["A21"].font,
                ws["B21"].value,
            ) = (
                "loc_td<X>d",
                FT_BOLD,
                "Локальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
            )
            (
                ws["A22"].value,
                ws["A22"].font,
                ws["B22"].value,
            ) = (
                "glob_td<X>d",
                FT_BOLD,
                "Глобальный прогноз на периоде обучения из <X> дней (time-depth=<X>d)",
            )

        elif data_type == "metrics_daily":
            (
                ws["A18"].value,
                ws["A18"].font,
                ws["B18"].value,
            ) = (
                "local",
                FT_BOLD,
                "Локальный прогноз вперед от указанной даты",
            )
            (
                ws["A19"].value,
                ws["A19"].font,
                ws["B19"].value,
            ) = (
                "global",
                FT_BOLD,
                "Глобальный прогноз вперед от указанной даты",
            )

            irow = 21
        elif data_type == "metrics_daily_sum":
            (
                ws["A18"].value,
                ws["A18"].font,
                ws["B18"].value,
            ) = (
                "local",
                FT_BOLD,
                "Локальный прогноз вперед",
            )
            (
                ws["A19"].value,
                ws["A19"].font,
                ws["B19"].value,
            ) = (
                "global",
                FT_BOLD,
                "Глобальный прогноз вперед",
            )

            ws["A21"] = "Метрики от метрик"
            ws["A21"].font = FT_BOLD_BLUE
            (
                ws["A22"].value,
                ws["A22"].font,
                ws["B22"].value,
            ) = (
                "min",
                FT_BOLD,
                "Минимальное значение метрики по серии прогнозов",
            )
            (
                ws["A23"].value,
                ws["A23"].font,
                ws["B23"].value,
            ) = (
                "max",
                FT_BOLD,
                "Максимальное значение метрики по серии прогнозов",
            )
            (
                ws["A24"].value,
                ws["A24"].font,
                ws["B24"].value,
            ) = (
                "mean",
                FT_BOLD,
                "Среднее значение метрики по серии прогнозов",
            )
            (
                ws["A25"].value,
                ws["A25"].font,
                ws["B25"].value,
            ) = (
                "std",
                FT_BOLD,
                "Стандартное отклонение метрики по серии прогнозов",
            )

            irow = 27

        if (data_type == "metrics_daily_sum") or (data_type == "metrics_daily"):
            ws.cell(column=1, row=irow).value = "Метрики"
            ws.cell(column=1, row=irow).font = FT_BOLD_BLUE

            (
                ws.cell(column=1, row=irow + 1).value,
                ws.cell(column=1, row=irow + 1).font,
                ws.cell(column=2, row=irow + 1).value,
            ) = ("min", FT_BOLD, "Минимальное значение")
            (
                ws.cell(column=1, row=irow + 2).value,
                ws.cell(column=1, row=irow + 2).font,
                ws.cell(column=2, row=irow + 2).value,
            ) = ("max", FT_BOLD, "Максимальное значение")
            (
                ws.cell(column=1, row=irow + 3).value,
                ws.cell(column=1, row=irow + 3).font,
                ws.cell(column=2, row=irow + 3).value,
            ) = ("mean", FT_BOLD, "Среднее значение")
            (
                ws.cell(column=1, row=irow + 4).value,
                ws.cell(column=1, row=irow + 4).font,
                ws.cell(column=2, row=irow + 4).value,
            ) = ("med", FT_BOLD, "Медианное значение")

            (
                ws.cell(column=1, row=irow + 5).value,
                ws.cell(column=1, row=irow + 5).font,
                ws.cell(column=2, row=irow + 5).value,
            ) = ("me", FT_BOLD, "Средняя ошибка")
            (
                ws.cell(column=1, row=irow + 6).value,
                ws.cell(column=1, row=irow + 6).font,
                ws.cell(column=2, row=irow + 6).value,
            ) = ("mae", FT_BOLD, "Средняя абсолютная ошибка")
            (
                ws.cell(column=1, row=irow + 7).value,
                ws.cell(column=1, row=irow + 7).font,
                ws.cell(column=2, row=irow + 7).value,
            ) = ("rmse", FT_BOLD, "Среднеквадратическая ошибка")
            (
                ws.cell(column=1, row=irow + 8).value,
                ws.cell(column=1, row=irow + 8).font,
                ws.cell(column=2, row=irow + 8).value,
            ) = ("r2", FT_BOLD, "Коэффициент детерминации")
            (
                ws.cell(column=1, row=irow + 9).value,
                ws.cell(column=1, row=irow + 9).font,
                ws.cell(column=2, row=irow + 9).value,
            ) = ("mre", FT_BOLD, "Средняя относительная ошибка (она же MAPE)")
            (
                ws.cell(column=1, row=irow + 10).value,
                ws.cell(column=1, row=irow + 10).font,
                ws.cell(column=2, row=irow + 10).value,
            ) = ("nse", FT_BOLD, "Коэффициент эффективности Нэша-Сатклиффа")
            (
                ws.cell(column=1, row=irow + 11).value,
                ws.cell(column=1, row=irow + 11).font,
                ws.cell(column=2, row=irow + 11).value,
            ) = ("kge", FT_BOLD, "Коэффициент эффективности Клинга-Гупты")

    def _write_ml_models_sheet_to_station_file(self, wb, dict_ml_models, st_code):
        logging.info(
            "Запись информации о ML-моделях в файл с метриками для станции %s",
            st_code,
        )

        ws = wb.create_sheet(title="ML-models", index=0)

        now_dates = sorted(list(dict_ml_models.keys()))

        columns = sorted(
            [par_name for par_name in dict_ml_models[now_dates[0]][st_code].keys()]
        )
        for idx, col_name in enumerate(columns):
            ws.cell(row=1, column=idx + 2).value = col_name
            ws.cell(row=1, column=idx + 2).font = FT_BOLD

        for idx in range(len(columns) + 1):
            ws.column_dimensions[get_column_letter(idx + 1)].width = COL_WIDHT_2X

        for idx_row, now_date in enumerate(now_dates):
            ws.cell(row=idx_row + 2, column=1).value = now_date
            for idx_col, par_name in enumerate(columns):
                ws.cell(row=idx_row + 2, column=idx_col + 2).value = dict_ml_models[
                    now_date
                ][st_code][par_name]

    def _write_ml_models_sheet_to_parameter_file(self, wb, dict_ml_models, par_name):
        logging.info(
            "Запись информации о ML-моделях в файл с метриками для параметра %s",
            par_name,
        )

        ws = wb.create_sheet(title="ML-models", index=0)

        now_dates = sorted(list(dict_ml_models.keys()))

        columns = sorted([st_code for st_code in dict_ml_models[now_dates[0]].keys()])

        for idx, col_name in enumerate(columns):
            ws.cell(row=1, column=idx + 2).value = col_name
            ws.cell(row=1, column=idx + 2).font = FT_BOLD

        for idx in range(len(columns) + 1):
            ws.column_dimensions[get_column_letter(idx + 1)].width = COL_WIDHT_2X

        for idx_row, now_date in enumerate(now_dates):
            ws.cell(row=idx_row + 2, column=1).value = now_date
            for idx_col, st_code in enumerate(columns):
                if par_name in dict_ml_models[now_date][st_code]:
                    ws.cell(row=idx_row + 2, column=idx_col + 2).value = dict_ml_models[
                        now_date
                    ][st_code][par_name]

    def _write_ml_models_sheet_to_expdata_file(
        self, wb, dict_ml_models, st_code, par_name
    ) -> None:
        logging.info(
            "Запись информации о ML-моделях в файл с экспортируемыми данными для станции %s и параметра %s",
            st_code,
            par_name,
        )

        ws = wb.create_sheet(title="ML-models", index=0)

        now_dates = sorted(list(dict_ml_models.keys()))

        ws.cell(row=1, column=1).value = "Дата"
        ws.cell(row=1, column=1).font = FT_BOLD

        for idx in range(2):
            ws.column_dimensions[get_column_letter(idx + 1)].width = COL_WIDHT_2X

        for idx_row, now_date in enumerate(now_dates):
            ws.cell(row=idx_row + 2, column=1).value = now_date
            if (
                st_code in dict_ml_models[now_date]
                and par_name in dict_ml_models[now_date][st_code]
            ):
                ws.cell(row=idx_row + 2, column=2).value = dict_ml_models[now_date][
                    st_code
                ][par_name]

    def _write_metrics_1_file_is_1_station_all_parameters(
        self, fname_prefix: str, info: dict, dfs: dict, dict_ml_models: dict
    ) -> None:
        logging.info(
            "Формирование серии файлов: 1 файл = 1 станция с метриками по всем метео-параметрам"
        )

        # Цикл по станциям:
        for station_code in dfs.keys():
            fname = fname_prefix + "_st-" + station_code + ".xlsx"
            logging.info("[%s] Запись файла %s", station_code, fname)

            out_file = os.path.join(self.output_folder, fname)

            # Создать файл и записать все датафреймы
            with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
                mm_list = []

                # Цикл по метео-параметрам:
                for par_name in sorted(dfs[station_code].keys()):
                    # Запись датафрейма dfs[station_code][par_name] на отдельный лист с именем par_name
                    dfs[station_code][par_name].to_excel(writer, sheet_name=par_name)

                    df = dfs[station_code][par_name].agg(["min", "max", "mean", "std"])
                    df.index = pd.MultiIndex.from_product([[par_name], df.index])
                    mm_list.append(df)

                mm_df = pd.concat(mm_list)
                mm_df.to_excel(writer, sheet_name="Metrics_summary")

            #  Открыть файл для дозаписи листа Info
            wb = load_workbook(out_file)

            # Переставить лист Metrics_summary на 1-ю позицию (индекс 0)
            metrics_summary_sheet = wb["Metrics_summary"]
            wb._sheets.remove(metrics_summary_sheet)
            wb._sheets.insert(0, metrics_summary_sheet)

            # Сформировать 1-й лист (ML-models), на основе словаря dict_ml_models (для записи используется openpyxl)
            self._write_ml_models_sheet_to_station_file(
                wb, dict_ml_models, station_code
            )

            # Сформировать 1-й лист (Info), на основе словаря info (для записи используется openpyxl)
            self._write_info_sheet_to_station_file(wb, info, station_code)

            wb.save(out_file)

    def _write_metrics_1_file_is_1_parameter_all_stations(
        self, fname_prefix: str, info: dict, dfs: dict, dict_ml_models: dict
    ) -> None:
        logging.info(
            "Формирование серии файлов: 1 файл = 1 метео-параметр (по имени) с метриками по станциям"
        )

        # Получить список всех метео-параметров из словаря dfs
        par_names = set()
        for station_code in dfs.keys():
            par_names.update(dfs[station_code].keys())
        par_names = sorted(list(par_names))

        # Цикл по именам метео-параметров из par_names:
        for par_name in par_names:
            fname = fname_prefix + "_par-" + par_name + ".xlsx"
            logging.info("[%s] Запись файла %s", par_name, fname)

            out_file = os.path.join(self.output_folder, fname)

            # Открыть файл с помощью pandas.ExcelWriter
            with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
                mm_list = []

                # Цикл по станциям:
                for station_code in sorted(dfs.keys()):
                    if par_name in dfs[station_code]:

                        # Запись датафрейма dfs[station_code][par_name] на отдельный лист с именем station_code
                        dfs[station_code][par_name].to_excel(
                            writer, sheet_name=station_code
                        )

                        df = dfs[station_code][par_name].agg(
                            ["min", "max", "mean", "std"]
                        )
                        df.index = pd.MultiIndex.from_product(
                            [[station_code], df.index]
                        )
                        mm_list.append(df)

                mm_df = pd.concat(mm_list)
                mm_df.to_excel(writer, sheet_name="Metrics_summary")

            #  Открыть файл для дозаписи листа Info
            wb = load_workbook(out_file)

            # Переставить лист Metrics_summary на 1-ю позицию (индекс 0)
            metrics_summary_sheet = wb["Metrics_summary"]
            wb._sheets.remove(metrics_summary_sheet)
            wb._sheets.insert(0, metrics_summary_sheet)

            # Сформировать 1-й лист (ML-models), на основе словаря dict_ml_models (для записи используется openpyxl)
            self._write_ml_models_sheet_to_parameter_file(wb, dict_ml_models, par_name)

            #   Сформировать 1-й лист (Info), на основе словаря info (для записи используется openpyxl)
            self._write_info_sheet_to_parameter_file(wb, info, par_name)

            wb.save(out_file)

    def _write_1_file_is_1_station_1_parameters(
        self,
        data_type: str,
        fname_prefix: str,
        info: dict,
        dfs: dict,
        dict_ml_models: dict,
    ) -> None:
        logging.info(
            "Формирование серии файлов: 1 файл = данные наблюдений и прогнозы для 1 метеопараметра 1 станции"
        )

        # Цикл по станциям:
        for station_code, station in dfs.items():
            # Цикл по метео-параметрам:
            for par_name, par_info in station.items():
                fname = (
                    fname_prefix + "_st-" + station_code + "_par-" + par_name + ".xlsx"
                )
                logging.info("[%s] [%s] Запись файла %s", station_code, par_name, fname)

                out_file = os.path.join(self.output_folder, fname)

                # Создать файл и записать датафреймы dfs[station_code][par_name]
                with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
                    for sheet_name, df in par_info.items():
                        df.to_excel(writer, sheet_name=sheet_name)

                #  Открыть файл для дозаписи листа Info
                wb = load_workbook(out_file)

                # Сформировать 1-й лист (ML-models), на основе словаря dict_ml_models (для записи используется openpyxl)
                self._write_ml_models_sheet_to_expdata_file(
                    wb, dict_ml_models, station_code, par_name
                )

                # Сформировать 1-й лист (Info), на основе словаря info (для записи используется openpyxl)
                self._write_info_sheet_to_expdata_file(
                    wb, data_type, info, station_code, par_name
                )

                wb.save(out_file)

    def write_metrics(self, in_data: dict, metrics: dict) -> None:
        logging.info("Запись вычисленных метрик в файлы")

        now_dates_list = list(metrics.keys())

        info = self._collect_info(in_data[list(in_data.keys())[0]])
        dict_ml_models = self._collect_ml_models_info(in_data, now_dates_list)

        dfs_metrics = self._gen_dfs_metrics(metrics)

        self._write_metrics_1_file_is_1_station_all_parameters(
            self.metrics_filename_prefix, info, dfs_metrics, dict_ml_models
        )
        self._write_metrics_1_file_is_1_parameter_all_stations(
            self.metrics_filename_prefix, info, dfs_metrics, dict_ml_models
        )

        dfs_metrics_daily, dfs_metrics_daily_sum = self._gen_dfs_metrics_daily(metrics)
        self._write_1_file_is_1_station_1_parameters(
            "metrics_daily",
            self.metrics_daily_filename_prefix,
            info,
            dfs_metrics_daily,
            dict_ml_models,
        )
        self._write_1_file_is_1_station_1_parameters(
            "metrics_daily_sum",
            self.metrics_daily_sum_filename_prefix,
            info,
            dfs_metrics_daily_sum,
            dict_ml_models,
        )

    def write_data(self, in_data: dict, exp_data: dict, now_dates: list) -> None:
        logging.info("Запись данных в файлы")

        info = self._collect_info(in_data[list(in_data.keys())[0]])
        dict_ml_models = self._collect_ml_models_info(in_data, now_dates)

        dfs_data = self._gen_dfs_data(exp_data, now_dates)

        self._write_1_file_is_1_station_1_parameters(
            "data_export", self.data_filename_prefix, info, dfs_data, dict_ml_models
        )
